import openpyxl
import sys

wb = openpyxl.load_workbook(r'C:\Users\Machcreator\Desktop\Robo Visual\case-01-robotic-solutions-platform-team-20\Датасеты_хакатон.xlsx', data_only=True)

with open('xl_output.txt', 'w', encoding='utf-8') as out:
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        out.write(f'\n--- {sheet_name} ---\n')
        for row in ws.iter_rows(min_row=1, max_row=min(100, ws.max_row), values_only=False):
            vals = [str(cell.value) if cell.value is not None else '' for cell in row]
            if any(v for v in vals):
                out.write(str(vals) + '\n')